import os
import tempfile
from app.processors.text_processor import text_processor
from app.processors.pdf_processor import pdf_processor
from app.processors.docx_processor import docx_processor
from app.processors.url_processor import url_processor

def test_text_processor():
    raw_text = "Critical Threat Alert: Ransomware targeting domain controllers."
    result = text_processor.process(raw_text, title="Alert Doc")
    assert result["source_type"] == "text"
    assert result["title"] == "Alert Doc"
    assert "Ransomware" in result["text"]
    assert len(result["pages"]) == 1

def test_docx_processor():
    try:
        import docx
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            doc_path = tmp.name
        
        doc = docx.Document()
        doc.add_heading("Threat Briefing", 0)
        doc.add_paragraph("CVE-2026-9999 is actively exploited in the wild.")
        doc.save(doc_path)

        result = docx_processor.process_file(doc_path)
        assert result["source_type"] == "docx"
        assert "CVE-2026-9999" in result["text"]

        if os.path.exists(doc_path):
            os.remove(doc_path)
    except ImportError:
        pass
