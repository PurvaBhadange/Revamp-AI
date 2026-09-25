import uuid
import os
from typing import Any, Dict

class DOCXProcessor:
    def process_file(self, file_path: str, title: str = None) -> Dict[str, Any]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found at {file_path}")

        doc_title = title or os.path.basename(file_path)
        paragraphs_text = []

        try:
            import docx
            doc = docx.Document(file_path)
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs_text.append(p.text.strip())
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs_text.append(row_text)
        except Exception as e:
            paragraphs_text = [f"Error extracting DOCX text: {str(e)}"]

        combined_text = "\n\n".join(paragraphs_text).strip()

        return {
            "document_id": str(uuid.uuid4()),
            "source_type": "docx",
            "title": doc_title,
            "text": combined_text,
            "language": "English",
            "metadata": {
                "file_path": file_path,
                "paragraph_count": len(paragraphs_text)
            },
            "pages": [{"page_number": 1, "text": combined_text}],
            "media": [],
            "entities": [],
            "timestamps": []
        }

docx_processor = DOCXProcessor()
